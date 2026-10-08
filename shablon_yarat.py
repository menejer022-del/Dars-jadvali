"""Kirish Excel shablonini namuna ma'lumotlar bilan yaratadi.

Ishlatish:  python3 shablon_yarat.py kirish.xlsx
Keyin kirish.xlsx ni o'z ma'lumotlaringiz bilan to'ldiring.
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

out = sys.argv[1] if len(sys.argv) > 1 else "kirish.xlsx"

HEAD = PatternFill("solid", fgColor="1F4E78")
NOTE = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="999999")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def sheet(wb, title, note, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.cell(1, 1, note).fill = NOTE
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(header))
    ws.cell(1, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 48
    for j, h in enumerate(header, 1):
        c = ws.cell(2, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    for i, r in enumerate(rows, 3):
        for j, v in enumerate(r, 1):
            c = ws.cell(i, j, v)
            c.border = BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + j)].width = w
    ws.freeze_panes = "A3"
    return ws


wb = Workbook()
wb.remove(wb.active)

# ---- Sozlamalar ----
ws = wb.create_sheet("Sozlamalar")
ws.cell(1, 1, "Kunlar (A-B ustun) va juftliklar (D-E ustun). Kod qisqa bo'lsin: Du, Se, ...").fill = NOTE
ws.merge_cells("A1:E1")
for col, h in zip("ABDE", ["Kod", "Kun nomi", "Juftlik", "Vaqti"]):
    c = ws[f"{col}2"]
    c.value = h
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = HEAD
days = [("Du", "Dushanba"), ("Se", "Seshanba"), ("Ch", "Chorshanba"),
        ("Pa", "Payshanba"), ("Ju", "Juma"), ("Sh", "Shanba")]
pairs = [(1, "8:00-9:20"), (2, "9:30-10:50"), (3, "11:00-12:20"), (4, "13:00-14:20")]
for i, (k, n) in enumerate(days, 3):
    ws.cell(i, 1, k)
    ws.cell(i, 2, n)
for i, (k, n) in enumerate(pairs, 3):
    ws.cell(i, 4, k)
    ws.cell(i, 5, n)
for col, w in zip("ABCDE", [8, 16, 3, 10, 16]):
    ws.column_dimensions[col].width = w

# ---- Guruhlar ----
G = [(f"Kiber {i}-25", 25, "Sh") for i in range(1, 7)]
sheet(wb, "Guruhlar",
      "Guruh nomi, talabalar soni va dars bo'lmaydigan kunlar (kod bilan, vergul orqali; masalan: Sh yoki Ch, Sh). "
      "Bo'sh kunlar jadvalda 'MUSTAQIL TA'LIM' deb chiqadi.",
      ["Guruh", "Talabalar soni", "Dars bo'lmaydigan kunlar"], G, [18, 16, 28])

# ---- Oqituvchilar ----
T = [
    ("Yakubbayeva N.", "Pa-3, Pa-4", 4),
    ("Mamatov A.", "", 4),
    ("Shabdarov E.", "", 4),
    ("Sharipov J.", "", 4),
    ("Sobirov R.", "Ju-1", 4),
    ("Qulmatov M.", "", 4),
    ("Aslamov A.", "", 4),
]
sheet(wb, "Oqituvchilar",
      "Band vaqtlar: dars bera olmaydigan vaqtlar 'Kun-Juftlik' ko'rinishida (Du-1, Pa-3) yoki butun kun (Se). "
      "Vergul bilan ajrating. Kuniga maksimal juftlik soni ham kiritiladi.",
      ["O'qituvchi", "Band vaqtlar", "Kuniga max juftlik"], T, [20, 30, 20])

# ---- Xonalar ----
R = [("211", 60), ("306", 120), ("401", 30), ("406", 30), ("407", 30),
     ("408", 30), ("411", 60), ("412", 30)]
sheet(wb, "Xonalar",
      "Xona raqami va sig'imi (o'rinlar soni). NAMUNA qiymatlar: haqiqiy sig'imlarni kiriting.",
      ["Xona", "Sig'im"], R, [12, 12])

# ---- Darslar ----
BD = "Boshlang'ich dasturlash va muammolarni yechish"
HA = "Hisoblash asoslari"
IA = "Internet dasturiy ta'minot arxitekturasi"
streams = ["Kiber 1-25, Kiber 2-25", "Kiber 3-25, Kiber 4-25", "Kiber 5-25, Kiber 6-25"]
L = []
# Boshlang'ich dasturlash
L += [(BD, "M", "Yakubbayeva N.", streams[0], 1), (BD, "M", "Yakubbayeva N.", streams[1], 1),
      (BD, "M", "Mamatov A.", streams[2], 1)]
for g in (1, 2, 3):
    L.append((BD, "A", "Shabdarov E.", f"Kiber {g}-25", 2))
for g in (4, 5, 6):
    L.append((BD, "A", "Sharipov J.", f"Kiber {g}-25", 2))
# Hisoblash asoslari
for s in streams:
    L.append((HA, "M", "Sobirov R.", s, 1))
for g in range(1, 7):
    L.append((HA, "A", "Sobirov R.", f"Kiber {g}-25", 1))
# Internet dasturiy ta'minot arxitekturasi
for s in streams:
    L.append((IA, "M", "Qulmatov M.", s, 1))
for g in (1, 2, 3, 4):
    L.append((IA, "A", "Qulmatov M.", f"Kiber {g}-25", 2))
for g in (5, 6):
    L.append((IA, "A", "Aslamov A.", f"Kiber {g}-25", 2))
L = [(i,) + r for i, r in enumerate(L, 1)]
sheet(wb, "Darslar",
      "Har qator = bitta fan/turi. 'Guruhlar' ustunida bir nechta guruh vergul bilan yozilsa, bu OQIM "
      "(ma'ruza birga o'tadi: bir vaqt, bir xona). 'Haftalik juftlik' = haftada necha juftlik.",
      ["№", "Fan", "Turi (M/A)", "O'qituvchi", "Guruhlar", "Haftalik juftlik"], L,
      [6, 46, 12, 20, 34, 16])

wb.save(out)
print("Shablon yaratildi:", out, "| darslar qatori:", len(L))
