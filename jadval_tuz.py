"""Dars jadvalini avtomatik tuzuvchi (Google OR-Tools CP-SAT).

Ishlatish:  python3 jadval_tuz.py kirish.xlsx natija.xlsx [vaqt_chegarasi_sekund]
"""
import os
import sys
from collections import defaultdict
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from ortools.sat.python import cp_model

# ------------------------------------------------------------------ o'qish


def split(s):
    return [x.strip() for x in str(s or "").replace(";", ",").split(",") if x.strip()]


def load(path):
    wb = load_workbook(path, data_only=True)
    ws = wb["Sozlamalar"]
    days, pairs = [], []
    for r in ws.iter_rows(min_row=3, values_only=True):
        if r[0]:
            days.append((str(r[0]).strip(), str(r[1]).strip()))
        if len(r) > 4 and r[3]:
            pairs.append((str(r[3]).strip(), str(r[4]).strip()))
    dcode = {c.lower(): i for i, (c, n) in enumerate(days)}
    dcode.update({n.lower(): i for i, (c, n) in enumerate(days)})
    pcode = {c: i for i, (c, _) in enumerate(pairs)}

    groups = {}  # nom -> (soni, bo'sh kunlar to'plami)
    for r in load_rows(wb["Guruhlar"]):
        free = set()
        for t in split(r[2]):
            if t.lower() not in dcode:
                raise SystemExit(f"Guruhlar: noma'lum kun '{t}'")
            free.add(dcode[t.lower()])
        groups[str(r[0]).strip()] = (int(r[1]), free)

    teachers = {}  # nom -> (band to'plam, kuniga max)
    for r in load_rows(wb["Oqituvchilar"]):
        busy = set()
        for t in split(r[1]):
            if "-" in t:
                dd, pp = [x.strip() for x in t.split("-", 1)]
                if dd.lower() not in dcode or pp not in pcode:
                    raise SystemExit(f"Oqituvchilar: noto'g'ri band vaqt '{t}'")
                busy.add((dcode[dd.lower()], pcode[pp]))
            else:
                if t.lower() not in dcode:
                    raise SystemExit(f"Oqituvchilar: noto'g'ri band kun '{t}'")
                for p in range(len(pairs)):
                    busy.add((dcode[t.lower()], p))
        teachers[str(r[0]).strip()] = (busy, int(r[2]) if r[2] else len(pairs))

    rooms = {str(r[0]).strip(): int(r[1]) for r in load_rows(wb["Xonalar"])}

    lessons = []
    for r in load_rows(wb["Darslar"]):
        gs = split(r[4])
        for g in gs:
            if g not in groups:
                raise SystemExit(f"Darslar: '{g}' guruhi Guruhlar varag'ida yo'q")
        t = str(r[3]).strip()
        teachers.setdefault(t, (set(), len(pairs)))
        lessons.append(dict(id=r[0], subject=str(r[1]).strip(), type=str(r[2]).strip(),
                            teacher=t, groups=gs, n=int(r[5])))
    return days, pairs, groups, teachers, rooms, lessons


def load_rows(ws):
    return [r for r in ws.iter_rows(min_row=3, values_only=True) if r[0] not in (None, "")]


# ------------------------------------------------------------------ model


def diagnose(data, events):
    """Model qurishdan oldin ma'lumotdagi aniq muammolarni topadi (o'zbekcha xabarlar bilan)."""
    days, pairs, groups, teachers, rooms, lessons = data
    D, P = len(days), len(pairs)
    issues = []

    t_n = defaultdict(int)
    g_n = defaultdict(int)
    for e in events:
        t_n[e["teacher"]] += 1
        for g in e["groups"]:
            g_n[g] += 1
        tb = teachers[e["teacher"]][0]
        ok = [1 for d in range(D) for p in range(P)
              if (d, p) not in tb and all(d not in groups[g][1] for g in e["groups"])]
        if not ok:
            issues.append(f"Dars №{e['id']} ({e['subject']}, {e['teacher']}): o'qituvchi va guruhlarning "
                          f"bo'sh vaqtlari umuman mos kelmaydi.")
    for t, n in t_n.items():
        busy, mx = teachers[t]
        free = D * P - len(busy)
        cap = min(free, mx * D)
        if n > cap:
            issues.append(f"O'qituvchi {t}: haftada {n} juftlik berilgan, lekin sig'adigan vaqt ko'pi bilan {cap} "
                          f"(band vaqtlar va kuniga {mx} juftlik chegarasi hisobga olinganda).")
    for g, n in g_n.items():
        free_days = D - len(groups[g][1])
        if n > free_days * P:
            issues.append(f"Guruh {g}: haftada {n} juftlik berilgan, lekin dars kunlari bo'yicha faqat "
                          f"{free_days * P} ta o'rin bor.")
    caps = sorted(rooms.values())
    stud = [e["students"] for e in events]
    for c in sorted(set(stud)):
        need = sum(1 for s in stud if s >= c)
        nrooms = sum(1 for cap in caps if cap >= c)
        if nrooms == 0:
            issues.append(f"{c} o'rinli xona yo'q: {need} ta dars shuncha o'rin talab qiladi (oqim ma'ruzalari "
                          f"talabalar sonini qo'shadi).")
        elif need > nrooms * D * P:
            issues.append(f"{c} va undan katta sig'imli xonalar yetmaydi: {need} ta juftlik kerak, "
                          f"xonalar ({nrooms} ta) haftada ko'pi bilan {nrooms * D * P} ta juftlik beradi.")
    return issues


def solve(data, tlimit, w_gap=10, w_tgap=3, w_late=1, w_over=3, w_same=5):
    days, pairs, groups, teachers, rooms, lessons = data
    D, P = len(days), len(pairs)

    events = []
    for L in lessons:
        for k in range(L["n"]):
            events.append(dict(L, k=k, eid=len(events),
                               students=sum(groups[g][0] for g in L["groups"])))

    problems = diagnose(data, events)
    if problems:
        raise SystemExit("Ma'lumotda muammo bor:\n- " + "\n- ".join(problems))

    m = cp_model.CpModel()
    x = {}  # (eid, d, p) -> BoolVar
    ev_slots = {}
    for e in events:
        tb = teachers[e["teacher"]][0]
        al = [(d, p) for d in range(D) for p in range(P)
              if (d, p) not in tb and all(d not in groups[g][1] for g in e["groups"])]
        ev_slots[e["eid"]] = al
        vs = []
        for (d, p) in al:
            v = m.NewBoolVar(f"x{e['eid']}_{d}_{p}")
            x[e["eid"], d, p] = v
            vs.append(v)
        m.AddExactlyOne(vs)

    # Xonalar: vaqt bo'yicha sig'im sinflari (xonani keyin aniq tayinlaymiz).
    # Har bir juftlikda, har bir talab darajasi (>= c o'rin) uchun: shunday darslar soni <= shunday xonalar soni.
    stud = {e["eid"]: e["students"] for e in events}
    caps = sorted(rooms.values())
    thresholds = sorted(set(stud.values()))
    slot_events = defaultdict(list)
    for (eid, d, p) in x:
        slot_events[d, p].append(eid)
    for (d, p), eids in slot_events.items():
        for c in thresholds:
            need = [x[eid, d, p] for eid in eids if stud[eid] >= c]
            nrooms = sum(1 for cap in caps if cap >= c)
            if len(need) > nrooms:
                m.Add(sum(need) <= nrooms)

    by_teacher = defaultdict(list)
    by_group = defaultdict(list)
    for e in events:
        by_teacher[e["teacher"]].append(e["eid"])
        for g in e["groups"]:
            by_group[g].append(e["eid"])

    def occupancy(eids, name):
        occ = {}
        for d in range(D):
            for p in range(P):
                terms = [x[eid, d, p] for eid in eids if (eid, d, p) in x]
                o = m.NewBoolVar(f"o_{name}_{d}_{p}")
                if terms:
                    m.Add(sum(terms) == o)  # bir vaqtda ko'pi bilan bitta dars
                else:
                    m.Add(o == 0)
                occ[d, p] = o
        return occ

    g_occ = {g: occupancy(eids, f"g{g}") for g, eids in by_group.items()}
    t_occ = {t: occupancy(eids, f"t{t}") for t, eids in by_teacher.items()}

    obj = []

    def gaps(occ, name, w):
        for d in range(D):
            a = m.NewBoolVar(f"a_{name}_{d}")
            first = m.NewIntVar(0, P - 1, f"f_{name}_{d}")
            last = m.NewIntVar(0, P - 1, f"l_{name}_{d}")
            cnt = sum(occ[d, p] for p in range(P))
            for p in range(P):
                m.AddImplication(occ[d, p], a)
                m.Add(first <= p).OnlyEnforceIf(occ[d, p])
                m.Add(last >= p).OnlyEnforceIf(occ[d, p])
            gap = m.NewIntVar(0, P, f"gap_{name}_{d}")
            m.Add(gap >= last - first + 1 - cnt - P * (1 - a))
            obj.append(w * gap)
            yield cnt

    for g, occ in g_occ.items():
        for d, cnt in zip(range(D), gaps(occ, f"G{g}", w_gap)):
            over = m.NewIntVar(0, P, f"over_{g}_{d}")
            m.Add(over >= cnt - 3)
            obj.append(w_over * over)
        for (d, p), o in occ.items():
            obj.append(w_late * p * o)

    for t, occ in t_occ.items():
        for d, cnt in zip(range(D), gaps(occ, f"T{t}", w_tgap)):
            m.Add(cnt <= teachers[t][1])

    by_lesson = defaultdict(list)
    for e in events:
        by_lesson[e["id"]].append(e["eid"])
    for lid, eids in by_lesson.items():
        if len(eids) < 2:
            continue
        idx = {eid: sum((d * P + p) * x[eid, d, p] for (d, p) in ev_slots[eid]) for eid in eids}
        for a, b in zip(eids, eids[1:]):
            m.Add(idx[a] + 1 <= idx[b])
        for d in range(D):
            c = sum(x[eid, d, p] for eid in eids for p in range(P) if (eid, d, p) in x)
            ex = m.NewIntVar(0, len(eids), f"ex_{lid}_{d}")
            m.Add(ex >= c - 1)
            obj.append(w_same * ex)

    m.Minimize(sum(obj))
    s = cp_model.CpSolver()
    s.parameters.max_time_in_seconds = tlimit
    s.parameters.num_workers = max(1, min(8, os.cpu_count() or 1))
    st = s.Solve(m)
    name = s.StatusName(st)
    print(f"Holat: {name}  | maqsad (kam=yaxshi): "
          f"{s.ObjectiveValue() if st in (cp_model.OPTIMAL, cp_model.FEASIBLE) else '-'}  "
          f"| vaqt: {s.WallTime():.1f}s | darslar (event): {len(events)}")
    if st not in (cp_model.OPTIMAL, cp_model.FEASIBLE):
        raise SystemExit("Jadval topilmadi: cheklovlar juda qattiq (o'qituvchi yuklamasi, band vaqtlar yoki "
                         "xonalar). Qidiruv vaqtini oshirib ko'ring yoki cheklovlarni yumshating.")

    # Xonalarni tayinlash: har bir juftlikda eng katta talabdan boshlab, eng kichik mos xonani beramiz.
    per_slot = defaultdict(list)
    for (eid, d, p), v in x.items():
        if s.Value(v):
            per_slot[d, p].append(eid)
    res = {}
    for (d, p), eids in per_slot.items():
        free = sorted(rooms.items(), key=lambda kv: (kv[1], kv[0]))
        for eid in sorted(eids, key=lambda i: -stud[i]):
            for j, (r, cap) in enumerate(free):
                if cap >= stud[eid]:
                    res[eid] = (d, p, r)
                    free.pop(j)
                    break
            else:
                raise SystemExit("Ichki xato: xona tayinlanmadi.")
    return events, res, name


# ------------------------------------------------------------------ tekshiruv


def verify(data, events, res):
    days, pairs, groups, teachers, rooms, lessons = data
    errs = []
    seen = defaultdict(list)
    for e in events:
        d, p, r = res[e["eid"]]
        seen[("xona", r, d, p)].append(e["eid"])
        seen[("oqit", e["teacher"], d, p)].append(e["eid"])
        for g in e["groups"]:
            seen[("guruh", g, d, p)].append(e["eid"])
            if d in groups[g][1]:
                errs.append(f"{g} bo'sh kunida dars bor")
        if (d, p) in teachers[e["teacher"]][0]:
            errs.append(f"{e['teacher']} band vaqtida dars bor")
        if rooms[r] < e["students"]:
            errs.append(f"{r} xona sig'imi yetmaydi")
    errs += [f"To'qnashuv: {k}" for k, v in seen.items() if len(v) > 1]
    return errs


# ------------------------------------------------------------------ chiqarish

thin = Side(style="thin", color="888888")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill("solid", fgColor="D9D9D9")
FREE = PatternFill("solid", fgColor="EDEDED")
STREAM = PatternFill("solid", fgColor="E2EFDA")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write(data, events, res, path, status):
    days, pairs, groups, teachers, rooms, lessons = data
    D, P = len(days), len(pairs)
    wb = Workbook()
    ws = wb.active
    ws.title = "Guruhlar jadvali"
    gnames = list(groups)
    cell_ev = {}
    for e in events:
        d, p, r = res[e["eid"]]
        for g in e["groups"]:
            cell_ev[g, d, p] = (e, r)

    for j, h in enumerate(["Kun", "Juftlik", "Vaqti"] + gnames, 1):
        c = ws.cell(1, j, h)
        c.font = Font(bold=True)
        c.fill = HEAD
        c.alignment = CENTER
        c.border = BORDER
    for d in range(D):
        for p in range(P):
            row = 2 + d * P + p
            ws.cell(row, 2, p + 1)
            ws.cell(row, 3, pairs[p][1])
            for j in (1, 2, 3):
                ws.cell(row, j).alignment = CENTER
                ws.cell(row, j).border = BORDER
            j = 0
            while j < len(gnames):
                g = gnames[j]
                col = 4 + j
                item = cell_ev.get((g, d, p))
                if item is None:
                    c = ws.cell(row, col, "MUSTAQIL TA'LIM" if d in groups[g][1] else None)
                    c.alignment = CENTER
                    c.border = BORDER
                    if d in groups[g][1]:
                        c.fill = FREE
                    j += 1
                    continue
                e, r = item
                k = j
                while k + 1 < len(gnames) and cell_ev.get((gnames[k + 1], d, p), (None,))[0] is e:
                    k += 1
                c = ws.cell(row, col, f"{e['subject']} ({e['type']})\n{e['teacher']}  {r}")
                c.alignment = CENTER
                if k > j:
                    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=4 + k)
                    c.fill = STREAM
                for cc in range(col, 4 + k + 1):
                    ws.cell(row, cc).border = BORDER
                j = k + 1
            ws.row_dimensions[row].height = 48
        ws.cell(2 + d * P, 1, days[d][1])
        ws.merge_cells(start_row=2 + d * P, start_column=1, end_row=1 + (d + 1) * P, end_column=1)
        ws.cell(2 + d * P, 1).alignment = Alignment(textRotation=90, horizontal="center", vertical="center")
        ws.cell(2 + d * P, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 13
    for j in range(len(gnames)):
        ws.column_dimensions[get_column_letter(4 + j)].width = 30
    ws.freeze_panes = "D2"

    # O'qituvchilar jadvali
    wt = wb.create_sheet("O'qituvchilar jadvali")
    tn = sorted({e["teacher"] for e in events})
    t_ev = {}
    for e in events:
        d, p, r = res[e["eid"]]
        t_ev[e["teacher"], d, p] = (e, r)
    for j, h in enumerate(["Kun", "Juftlik", "Vaqti"] + tn, 1):
        c = wt.cell(1, j, h)
        c.font = Font(bold=True)
        c.fill = HEAD
        c.alignment = CENTER
        c.border = BORDER
    for d in range(D):
        for p in range(P):
            row = 2 + d * P + p
            for j, v in enumerate([days[d][1], p + 1, pairs[p][1]], 1):
                c = wt.cell(row, j, v)
                c.alignment = CENTER
                c.border = BORDER
            for j, t in enumerate(tn):
                c = wt.cell(row, 4 + j)
                c.border = BORDER
                c.alignment = CENTER
                if (t, d, p) in t_ev:
                    e, r = t_ev[t, d, p]
                    c.value = f"{e['subject']} ({e['type']})\n{', '.join(e['groups'])}  {r}"
            wt.row_dimensions[row].height = 48
    for col, w in zip("ABC", [12, 8, 13]):
        wt.column_dimensions[col].width = w
    for j in range(len(tn)):
        wt.column_dimensions[get_column_letter(4 + j)].width = 30
    wt.freeze_panes = "D2"

    # Hisobot
    wr = wb.create_sheet("Hisobot")
    wr.append(["Holat", status])
    wr.append([])
    wr.append(["Guruh", "Haftalik juftliklar", "Oynalar (bo'sh juftlik)", "Eng ko'p kunlik yuklama"])
    for c in wr[3]:
        c.font = Font(bold=True)
    for g in gnames:
        total = gaps_n = mx = 0
        for d in range(D):
            ps = sorted(p for p in range(P) if (g, d, p) in cell_ev)
            total += len(ps)
            mx = max(mx, len(ps))
            if ps:
                gaps_n += ps[-1] - ps[0] + 1 - len(ps)
        wr.append([g, total, gaps_n, mx])
    wr.append([])
    wr.append(["O'qituvchi", "Haftalik juftliklar", "Oynalar", "Eng ko'p kunlik yuklama"])
    for c in wr[wr.max_row]:
        c.font = Font(bold=True)
    for t in tn:
        total = gaps_n = mx = 0
        for d in range(D):
            ps = sorted(p for p in range(P) if (t, d, p) in t_ev)
            total += len(ps)
            mx = max(mx, len(ps))
            if ps:
                gaps_n += ps[-1] - ps[0] + 1 - len(ps)
        wr.append([t, total, gaps_n, mx])
    for col, w in zip("ABCD", [22, 20, 24, 26]):
        wr.column_dimensions[col].width = w
    wb.save(path)


def tuz_bytes(file, tlimit=60, **weights):
    """file: yo'l yoki fayl-obyekt. Qaytaradi: (natija_bytes, holat, xatolar, (data, events, res))."""
    import io
    data = load(file)
    events, res, status = solve(data, tlimit, **weights)
    errs = verify(data, events, res)
    buf = io.BytesIO()
    write(data, events, res, buf, status)
    return buf.getvalue(), status, errs, (data, events, res)


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else "kirish.xlsx"
    dst = sys.argv[2] if len(sys.argv) > 2 else "natija.xlsx"
    tl = float(sys.argv[3]) if len(sys.argv) > 3 else 60
    data = load(src)
    events, res, status = solve(data, tl)
    errs = verify(data, events, res)
    print("Mustaqil tekshiruv:", "XATO yo'q" if not errs else errs)
    write(data, events, res, dst, status)
    print("Saqlandi:", dst)
